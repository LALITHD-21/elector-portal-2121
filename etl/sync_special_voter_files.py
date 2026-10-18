#!/usr/bin/env python3
r"""
Special Voter Files Synchronization Pipeline
Parses the 9 special roll files directly in voter/:
10e, 112, 13E, 21, 28E, 32(1), 36e(1), 72e, EN_Form18_SE_2_WithPhoto
Enforces robust deduplication on epic_number and deterministic ID generation for blank EPICs.
"""

import os
import re
import sys
import json
import time
import ssl
import openpyxl
import urllib.request
import urllib.parse

SUPABASE_URL = "https://fxnomgqqenaqpuzervxd.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImZ4bm9tZ3FxZW5hcXB1emVydnhkIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc5MjMzNzA5MywiZXhwIjoyMTA3OTEzMDkzfQ.c67TOlHqgzjRQu2uvmhJngSFyeJQAQ6s5n_utqPosmM"

BASE_DIR = r"C:\Users\techb\Desktop\electrol-portal-2121\electrol-portal-NEW-main"
VOTER_DIR = os.path.join(BASE_DIR, "voter")

with open(os.path.join(BASE_DIR, "etl", "booths_master.json"), "r", encoding="utf-8") as f:
    BOOTH_MASTER = json.load(f)

with open(os.path.join(BASE_DIR, "etl", "polling_stations_master.json"), "r", encoding="utf-8") as f:
    POLLING_MASTER = json.load(f)

def resolve_polling_info(part_str, serial_num):
    part_key = str(part_str).strip()
    candidates = [p for p in POLLING_MASTER if str(p.get("basePart")).strip() == part_key]
    if not candidates:
        return None, None
    if serial_num is not None:
        for c in candidates:
            min_s = c.get("minSerial") or 1
            max_s = c.get("maxSerial") or 999999
            if min_s <= serial_num <= max_s:
                return c.get("buildingName"), c.get("pollingArea")
    primary = candidates[0]
    return primary.get("buildingName"), primary.get("pollingArea")

def clean_epic(val, part=None, serial=None, idx=0):
    if not val:
        return f"P{part}_SP_{idx:05d}"
    s = str(val).strip().upper()
    s = re.sub(r'[\s\.\-_/\\,]', '', s)
    if len(s) > 16 or any(w in s for w in ['SCHOOL', 'ROOM', 'MAIDAN', 'HOSPITAL', 'COLLEGE', 'BUILDING', 'PANCHAYAT', 'TEMPLE', 'GOVT']):
        return f"P{part}_SP_{idx:05d}"
    if len(s) < 4:
        return f"P{part}_SP_{idx:05d}"
    if s.startswith("NM0"): s = "NMO" + s[3:]
    elif s.startswith("1KI"): s = "IKI" + s[3:]
    elif s.startswith("5GZ"): s = "SGZ" + s[3:]
    return s

def clean_serial(val):
    if val is None: return None
    s = str(val).strip()
    s = re.sub(r'[#\s\r\n]', '', s)
    m = re.search(r'\d+', s)
    if m:
        try: return int(m.group(0))
        except: pass
    return None

def clean_text(val):
    if val is None: return None
    t = str(val).strip()
    t = re.sub(r'\s+', ' ', t)
    t = re.sub(r'^[\.\-,:;]+|[\.\-,:;]+$', '', t).strip()
    return t if t else None

def clean_sex(val):
    if not val: return None
    s = str(val).strip().upper()
    if 'M' in s or s.startswith('PUR') or s.startswith('MALE'): return 'M'
    if 'F' in s or s.startswith('STR') or s.startswith('FEMALE') or s.startswith('W'): return 'F'
    return None

def clean_age(val):
    if val is None: return None
    try:
        m = re.search(r'\d+', str(val))
        if m:
            age = int(m.group(0))
            if 18 <= age <= 125: return age
    except: pass
    return None

def post_records(chunk):
    url = f"{SUPABASE_URL}/rest/v1/electors?on_conflict=epic_number"
    headers = {
        'apikey': SUPABASE_KEY,
        'Authorization': f'Bearer {SUPABASE_KEY}',
        'Content-Type': 'application/json',
        'Prefer': 'resolution=merge-duplicates,return=minimal'
    }
    ctx = ssl.create_default_context()
    data = json.dumps(chunk).encode('utf-8')
    req = urllib.request.Request(url, data=data, headers=headers, method='POST')
    with urllib.request.urlopen(req, context=ctx, timeout=60) as resp:
        return len(chunk)

def upsert_batch(records, batch_size=200):
    total = len(records)
    upserted = 0
    for i in range(0, total, batch_size):
        chunk = records[i:i + batch_size]
        success = False
        for retry in range(2):
            try:
                upserted += post_records(chunk)
                success = True
                break
            except Exception as e:
                time.sleep(1)
        if not success:
            micro_size = 25
            for j in range(0, len(chunk), micro_size):
                sub_chunk = chunk[j:j + micro_size]
                try:
                    upserted += post_records(sub_chunk)
                except Exception:
                    for item in sub_chunk:
                        try:
                            upserted += post_records([item])
                        except Exception: pass
    return upserted

def get_db_total():
    try:
        url = f"{SUPABASE_URL}/rest/v1/electors?select=count"
        headers = {
            'apikey': SUPABASE_KEY,
            'Authorization': f'Bearer {SUPABASE_KEY}',
            'Prefer': 'count=exact'
        }
        ctx = ssl.create_default_context()
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, context=ctx, timeout=30) as resp:
            cr = resp.headers.get('content-range')
            if cr and '/' in cr: return int(cr.split('/')[-1])
    except Exception as e:
        print(f"Error getting db total: {e}")
    return None

def parse_special_file(fpath):
    fname = os.path.basename(fpath)
    wb = openpyxl.load_workbook(fpath, data_only=True)
    ws = wb.active
    
    # Extract Part Number
    m = re.match(r'^(\d+)', fname)
    part = m.group(1) if m else ('2' if 'Form18_SE_2' in fname else None)
    if not part:
        wb.close()
        return []

    booth_info = BOOTH_MASTER.get(part, {})
    district = booth_info.get("district")
    taluk = booth_info.get("taluk")
    ac_name = booth_info.get("ac_name")

    # Find Header Row
    h_idx = None
    col_map = {}
    for r_i, r in enumerate(ws.iter_rows(min_row=1, max_row=25, values_only=True), 1):
        txt = ' '.join([str(c) for c in r if c is not None]).lower()
        if 'name of the elector' in txt:
            h_idx = r_i
            for c_i, cell in enumerate(r):
                if not cell: continue
                ct = str(cell).strip().lower()
                if any(k in ct for k in ['sino', 'sl no', 'slno', 'serial']): col_map.setdefault('serial', c_i)
                elif 'name of the elector' in ct: col_map.setdefault('name', c_i)
                elif any(k in ct for k in ['father', 'mother', 'husband', 'relative']): col_map.setdefault('relative', c_i)
                elif 'address' in ct: col_map.setdefault('address', c_i)
                elif 'qualification' in ct: col_map.setdefault('qualification', c_i)
                elif 'occupation' in ct: col_map.setdefault('occupation', c_i)
                elif 'age' in ct: col_map.setdefault('age', c_i)
                elif 'sex' in ct or 'gender' in ct: col_map.setdefault('sex', c_i)
                elif 'epic' in ct: col_map.setdefault('epic', c_i)
            break

    records_map = {}
    current_serial = 0
    row_idx = 0
    in_deletions = False

    for r in ws.iter_rows(min_row=1, values_only=True):
        if not r or not any(r): continue
        txt = ' '.join([str(c) for c in r if c is not None]).lower()

        # Track deletion vs addition supplement sections
        if 'deletions list' in txt or 'iii. deletions' in txt:
            in_deletions = True
            continue
        if 'additions list' in txt or 'ii. additions' in txt:
            in_deletions = False
            continue
        if in_deletions:
            continue

        # Dynamic sub-header detection for supplements
        if 'name of the elector' in txt:
            col_map = {}
            for c_i, cell in enumerate(r):
                if not cell: continue
                ct = str(cell).strip().lower()
                if any(k in ct for k in ['sino', 'sl no', 'slno', 'serial']): col_map.setdefault('serial', c_i)
                elif 'name of the elector' in ct: col_map.setdefault('name', c_i)
                elif any(k in ct for k in ['father', 'mother', 'husband', 'relative']): col_map.setdefault('relative', c_i)
                elif 'address' in ct: col_map.setdefault('address', c_i)
                elif 'qualification' in ct: col_map.setdefault('qualification', c_i)
                elif 'occupation' in ct: col_map.setdefault('occupation', c_i)
                elif 'age' in ct: col_map.setdefault('age', c_i)
                elif 'sex' in ct or 'gender' in ct: col_map.setdefault('sex', c_i)
                elif 'epic' in ct: col_map.setdefault('epic', c_i)
            continue

        if not col_map or 'name' not in col_map:
            continue

        name_c = col_map.get('name')
        name = clean_text(r[name_c]) if (name_c is not None and name_c < len(r)) else None
        if not name or name.lower() in ['name', 'name of the elector', 'elector name']:
            continue

        row_idx += 1

        serial_c = col_map.get('serial')
        raw_serial = r[serial_c] if (serial_c is not None and serial_c < len(r)) else None
        serial = clean_serial(raw_serial)
        if serial is not None: current_serial = serial
        else:
            current_serial += 1
            serial = current_serial

        epic_c = col_map.get('epic')
        raw_epic = r[epic_c] if (epic_c is not None and epic_c < len(r)) else None
        epic = clean_epic(raw_epic, part=part, serial=serial, idx=row_idx)

        rel_c = col_map.get('relative')
        relative_name = clean_text(r[rel_c]) if (rel_c is not None and rel_c < len(r)) else None

        addr_c = col_map.get('address')
        address = clean_text(r[addr_c]) if (addr_c is not None and addr_c < len(r)) else None

        qual_c = col_map.get('qualification')
        qualification = clean_text(r[qual_c]) if (qual_c is not None and qual_c < len(r)) else None

        occ_c = col_map.get('occupation')
        occupation = clean_text(r[occ_c]) if (occ_c is not None and occ_c < len(r)) else None

        age_c = col_map.get('age')
        age = clean_age(r[age_c]) if (age_c is not None and age_c < len(r)) else None

        sex_c = col_map.get('sex')
        sex = clean_sex(r[sex_c]) if (sex_c is not None and sex_c < len(r)) else None

        b_name, b_area = resolve_polling_info(part, serial)

        records_map[epic] = {
            "part_number": part,
            "serial_number": serial,
            "epic_number": epic,
            "name": name,
            "relative_name": relative_name,
            "age": age,
            "sex": sex,
            "address": address,
            "qualification": qualification,
            "occupation": occupation,
            "district": district,
            "ac_name": ac_name,
            "taluk": taluk,
            "polling_station_name": b_name,
            "polling_address": b_area,
            "source_file": fname
        }

    wb.close()
    return list(records_map.values())

def main():
    print("=" * 70)
    print("SPECIAL VOTER FILES INGESTION PIPELINE")
    print("=" * 70)

    start_db_total = get_db_total()
    print(f"Current Electors in Supabase before sync: {start_db_total:,}")

    special_files = [
        '10e_1791811086.xlsx', '112_1791587543.xlsx', '13E_1791811086.xlsx',
        '21_1791587025.xlsx', '28E_1791811630.xlsx', '32_1791587153 (1).xlsx',
        '36e_1791966609 (1).xlsx', '72e_1791875014.xlsx',
        'EN_Form18_SE_2_WithPhoto (1)_1792218075.xlsx'
    ]

    total_parsed = 0
    total_upserted = 0

    for idx, fname in enumerate(special_files, 1):
        fpath = os.path.join(VOTER_DIR, fname)
        if not os.path.exists(fpath):
            print(f"[{idx}/{len(special_files)}] Skipping missing {fname}")
            continue

        records = parse_special_file(fpath)
        if not records:
            print(f"[{idx}/{len(special_files)}] 0 records from {fname}")
            continue

        part = records[0]["part_number"]
        total_parsed += len(records)

        upserted = upsert_batch(records, batch_size=200)
        total_upserted += upserted

        print(f"[{idx}/{len(special_files)}] Part {part:3s} | {fname[:38]:38s} | Parsed: {len(records):5d} | Upserted: {upserted:5d}")

    end_db_total = get_db_total()

    print("\n" + "=" * 70)
    print("SPECIAL FILES INGESTION SUMMARY")
    print("=" * 70)
    print(f"Total Unique Electors Extracted : {total_parsed:,}")
    print(f"Total Records Synced            : {total_upserted:,}")
    print(f"Starting Database Total         : {start_db_total:,}")
    print(f"Final Database Total            : {end_db_total:,} (Net Added: {end_db_total - (start_db_total or 0):,})")
    print(f"Official Target                 : 183,893")
    print("=" * 70)

if __name__ == "__main__":
    main()
